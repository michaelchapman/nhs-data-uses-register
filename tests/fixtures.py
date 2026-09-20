"""A small synthetic register, so tests never depend on the real data.

Three agreements chosen to exercise what has gone wrong before: a dataset the
register renamed part-way through (named one way on one version, another way on
the next, with file releases under both), a name with a comma in it, and an
organisation that appears only as a joint controller.
"""

from __future__ import annotations

import io
import json
import tempfile
from contextlib import contextmanager
from pathlib import Path
from unittest import mock

import openpyxl

from pipeline import aliases

OLD_NAME = "Maternity Services Data Set v1.5"
NEW_NAME = "MSDS (Maternity Services Data Set) v1.5"

AGREEMENT_HEADER = [
    "Reference Number", "Application Title", "Applicant Organisation",
    "Applicant Organisation Type", "Data Controller(s)", "Sole/Joint Data Controller",
    "DSA Start Date", "DSA End Date", "Does Sublicensing Apply", "For Commercial Purposes",
    "Objective for Processing", "Processing Activities", "Expected Output",
    "Expected Measurable Benefits", "Yielded Benefits",
]
DATASET_HEADER = [
    "Reference Number", "Dataset", "Type of Data", "Sensitive or Non-Sensitive", "Frequency",
    "Legal Basis for Provision of Data", "Common Law Duty of Confidentiality",
]
RELEASE_HEADER = ["Reference Number", "Dataset", "Month File Released", "Patient Opt-Outs Applied"]


def _agreement(reference, title, organisation, controllers, start, end):
    return [
        reference, title, organisation, "Academic", controllers,
        "Joint Data Controller" if ";" in controllers else "Sole Data Controller",
        start, end, "No", "No", "Objective", "Activities", "Output", "Benefits", "",
    ]


def _dataset(reference, name):
    return [reference, name, "Identifiable", "Sensitive", "One-off", "Consent", "Yes"]


def workbook_bytes() -> bytes:
    book = openpyxl.Workbook()
    agreements = book.active
    agreements.title = "Agreements"
    agreements.append(AGREEMENT_HEADER)
    agreements.append(_agreement(
        "DARS-NIC-1-AAAAA-v1", "Maternity study", "UNIVERSITY OF EXAMPLE",
        "UNIVERSITY OF EXAMPLE", "2020-01-01", "2021-01-01"))
    agreements.append(_agreement(
        "DARS-NIC-1-AAAAA-v2", "Maternity study", "UNIVERSITY OF EXAMPLE",
        "UNIVERSITY OF EXAMPLE; OTHER TRUST", "2021-01-01", "2030-01-01"))
    agreements.append(_agreement(
        "DARS-NIC-2-BBBBB-v1", "Ambulance study",
        "NHS Bristol, North Somerset and South Gloucestershire ICB - 15C",
        "NHS Bristol, North Somerset and South Gloucestershire ICB - 15C",
        "2022-01-01", "2031-01-01"))

    datasets = book.create_sheet("Datasets")
    datasets.append(DATASET_HEADER)
    datasets.append(_dataset("DARS-NIC-1-AAAAA-v1", OLD_NAME))
    datasets.append(_dataset("DARS-NIC-1-AAAAA-v2", NEW_NAME))
    datasets.append(_dataset("DARS-NIC-2-BBBBB-v1", NEW_NAME))
    datasets.append(_dataset("DARS-NIC-2-BBBBB-v1", "Other Data Set"))

    releases = book.create_sheet("DataReleases")
    releases.append(RELEASE_HEADER)
    for month in ("Jan-20", "Feb-20", "Mar-20"):
        releases.append(["DARS-NIC-1-AAAAA-v1", OLD_NAME, month, "Yes"])
    for month in ("Jan-22", "Feb-22"):
        releases.append(["DARS-NIC-1-AAAAA-v2", NEW_NAME, month, "Yes"])
    releases.append(["DARS-NIC-2-BBBBB-v1", NEW_NAME, "Jun-22", "No"])

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


@contextmanager
def dataset_aliases():
    """Point the alias files at temporary ones that merge the renamed dataset."""
    with tempfile.TemporaryDirectory() as directory:
        directory = Path(directory)
        datasets = directory / "dataset-aliases.json"
        datasets.write_text(json.dumps({
            "aliases": [{"canonical": NEW_NAME, "variants": [NEW_NAME, OLD_NAME]}],
            "ignored": [],
        }))
        organisations = directory / "organisation-aliases.json"
        organisations.write_text(json.dumps({"aliases": [], "ignored": []}))
        with mock.patch.object(aliases, "DATASET_ALIASES_PATH", datasets), \
                mock.patch.object(aliases, "ALIASES_PATH", organisations):
            yield


def site_meta(base_path: str = "") -> dict:
    return {
        "site_name": "Test site",
        "site_url": "https://example.test",
        "base_path": base_path,
        "repo_url": "https://example.test/repo",
        "register": "data-uses-register",
        "register_name": "Data Uses Register",
        "edition": "september2026",
        "retrieved": "2026-09-20T00:00:00+00:00",
        "source_url": "https://example.test/source.xlsx",
        "source_file": "source.xlsx",
        "source_page": "https://example.test/register",
        "archive_page": "https://example.test/archive",
        "today": "2026-09-20",
        "editions": [
            {"edition": "september2026", "retrieved": "2026-09-20", "counts": {"agreement_versions": 3}}
        ],
    }


FIRST_EDITION = {
    "comparable": False, "reason": "first-edition", "previous_edition": None,
    "added": [], "amended": [], "removed": [],
}
